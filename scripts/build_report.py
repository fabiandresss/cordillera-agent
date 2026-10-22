"""Genera el informe técnico (Word, máx. 5 páginas) en docs/informe/Informe_EP2_ISY0101.docx.

Las secciones que la evaluación obliga a redactar al equipo SIN apoyo de IA (justificaciones,
conclusiones y reflexiones individuales) quedan como recuadros por completar.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "informe"
OUT.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin = sec.bottom_margin = Cm(1.7)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.0
for name, size, color in (("Heading 1", 13, "1F3A5F"), ("Heading 2", 11, "2F5597")):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = RGBColor.from_string(color)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def P(text="", bold=False, italic=False, size=None, align=None, after=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rich(parts, style=None):
    p = doc.add_paragraph(style=style)
    for text, b in parts:
        p.add_run(text).bold = b
    return p


def bullet(parts):
    p = rich(parts if isinstance(parts, list) else [(parts, False)], style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    return p


def table(rows, widths, header=True, size=8.5, fill="DCE6F2"):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            if header and i == 0:
                r.bold = True
                shade(c, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def box(title, lines, fill="FFF4CE"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, fill)
    c.text = ""
    p = c.paragraphs[0]
    p.add_run(title).bold = True
    p.paragraph_format.space_after = Pt(1)
    for ln in lines:
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(1)
        r = q.add_run(ln)
        r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def img(path, width_cm, caption):
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    P(caption, italic=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, color="555555")


# ------------------------------------------------------------------ portada
P("Evaluación Parcial N°2 · ISY0101 Ingeniería de Soluciones con IA", size=9, color="555555", after=0)
P("Agente de despacho y rutas «Transportes Cordillera»", bold=True, size=18, color="1F3A5F", after=0)
P("Agente funcional con LangGraph, memoria corta y larga, y planificación adaptativa (stack 100% gratuito)",
  italic=True, size=10, after=3)
rich([("Integrantes: ", True), ("[Nombre Apellido 1] · [Nombre Apellido 2]      ", False),
      ("Docente: ", True), ("[Nombre]      ", False), ("Fecha: ", True), ("[dd/mm/aaaa]", False)])
rich([("Repositorio: ", True), ("[https://github.com/<usuario>/cordillera-dispatch-agent]", False)])

# ------------------------------------------------------------------ 1
doc.add_heading("1. Contexto y requerimiento organizacional", 1)
P("Transportes Cordillera (empresa ficticia) asigna a mano los pedidos del día a sus 5 camiones. Cuando un "
  "camión se avería, un cliente VIP pide urgencia o una entrega falla, el plan se descuadra, se pierde el "
  "contexto de lo ocurrido y los clientes no reciben avisos a tiempo. Se desarrolló un agente que recibe "
  "eventos en lenguaje natural y responde planificando, replanificando, avisando y registrando, con "
  "escalamiento a un supervisor humano cuando corresponde.")

# ------------------------------------------------------------------ 2
doc.add_heading("2. Arquitectura y orquestación de componentes (IE2, IE6)", 1)
img(ROOT / "docs" / "diagramas" / "arquitectura.png", 15.2,
    "Figura 1. Componentes del agente: interfaz, grafo LangGraph, herramientas, almacenes y LLM.")
P("El flujo se implementó como un grafo de estado en LangGraph (13 nodos, 4 aristas condicionales). La "
  "decisión crítica (prioridad, capacidad, ventanas) la toma un planificador determinista; el LLM "
  "interpreta texto libre y redacta avisos. El proveedor de LLM es intercambiable (Groq, Ollama o modo "
  "sin LLM) mediante una variable de entorno.")
img(ROOT / "docs" / "diagramas" / "flujo_grafo.png", 16.0,
    "Figura 2. Flujo del grafo: clasificar → memoria → actuar → planificar → validar/reparar → ejecutar → avisar.")

P("El agente se expone mediante una API HTTP (FastAPI) y un panel web en React construido con el kit de "
  "componentes fintech-ui-react (tema indigo): centro de control con KPIs, consola de eventos en lenguaje "
  "natural, razonamiento paso a paso, flota, rutas, pedidos, memoria con búsqueda semántica y asistente "
  "conversacional. Es responsivo (sidebar en escritorio, dock en móvil).")
img(ROOT / "docs" / "capturas" / "dashboard.jpg", 13.0,
    "Figura 3. Panel web: tras una avería (camión 3) y un pedido VIP urgente, el agente muestra su razonamiento.")

# ------------------------------------------------------------------ 3
doc.add_heading("3. Herramientas del agente (IE1)", 1)
table([
    ["Tipo", "Herramientas (src/tools/)", "Función"],
    ["Consulta", "get_fleet_status, get_pending_orders, get_traffic, get_client_profile, get_route, "
                 "search_past_incidents, search_policies", "Leen flota, pedidos, tráfico simulado, historial del cliente y memoria semántica"],
    ["Escritura", "assign_orders, write_route_sheet, notify_client, notify_supervisor, reschedule_order, "
                  "set_order_priority, set_truck_status, log_incident", "Modifican la BD y generan hojas de ruta (.md) y avisos (.jsonl); assign_orders valida capacidad"],
    ["Razonamiento", "classify_event (LLM + reglas), redacción de avisos, agente ReAct de consultas (qa_agent.py)",
     "Clasifica eventos con salida estructurada y validación de ids; el LLM elige herramientas de consulta"],
], [2.0, 8.2, 7.6])
P("Autonomía y límites: el agente ejecuta sin intervención el flujo completo; ante datos faltantes pregunta, "
  "ante eventos no reconocidos escala, y nunca cancela pedidos por su cuenta.")

# ------------------------------------------------------------------ 4
doc.add_heading("4. Memoria y recuperación de contexto (IE3, IE4)", 1)
bullet([("Memoria corta: ", True), ("checkpointer SQLite de LangGraph con un thread_id por jornada. Guarda el "
        "historial de eventos y el número de replanificaciones; sobrevive al reinicio del proceso (prueba "
        "test_s5b) y se usa para decidir: 3 replanes en la jornada generan un aviso de inestabilidad.", False)])
bullet([("Memoria larga semántica: ", True), ("ChromaDB (distancia coseno) con embeddings nomic-embed-text. "
        "Indexa 15 incidentes históricos y las 7 políticas; cada incidente nuevo se indexa al registrarse. "
        "Recupera la estrategia usada antes (p. ej. reprogramar mañana si el cliente tiene ausencias previas).", False)])
table([
    ["Consulta de prueba", "Mejor similitud"],
    ["«el camión 3 se averió en zona norte»", "0.717"],
    ["«se me rompió una rueda del camión» (paráfrasis)", "0.667"],
    ["«Cliente ausente en el pedido 8»", "0.739"],
    ["«pasta de dientes» (ruido)", "0.540 → descartado (umbral 0.60)"],
], [10.5, 7.3])

# ------------------------------------------------------------------ 5
doc.add_heading("5. Planificación y toma de decisiones (IE5)", 1)
P("Puntuación: score = prioridad×100 + 50 (si VIP) + 100/max(holgura, 5). Asignación voraz por score con "
  "capacidad como restricción dura y bono de estabilidad; secuenciación por cierre de ventana, viaje y "
  "prioridad; validación (CAPACIDAD, VENTANA, TURNO) y reparación local hasta 3 intentos, tras los cuales "
  "se escala. Decisiones por evento:")
table([
    ["Evento", "Decisión del agente"],
    ["Avería de camión", "Camión → averiado; reasigna sus pedidos pendientes entre los demás; avisa a los afectados"],
    ["VIP urgente", "Prioridad 4 y hora límite; reordena la ruta; si no es factible, repara y escala"],
    ["Entrega fallida", "1.er intento: reintento hoy (o mañana si el cliente tiene ausencias previas); 2.º intento: escala al supervisor"],
    ["Ambiguo / desconocido", "Pide el dato faltante o escala; descarta ids inexistentes"],
], [4.0, 13.8])

# ------------------------------------------------------------------ 6
doc.add_heading("6. Pruebas y evidencia", 1)
P("Se implementaron 57 pruebas automáticas (pytest, sin red ni LLM) y 9 escenarios con registro del "
  "razonamiento nodo por nodo (docs/evidencia/).")
table([
    ["Escenario", "Resultado verificado"],
    ["1 Avería (11:00)", "9 entregas ya realizadas; 16 pendientes reasignadas; 0 pedidos en el camión averiado"],
    ["2 VIP urgente", "Pedido 14 con prioridad 4, 1.ª parada de su camión y ETA 12:00 (límite 12:00)"],
    ["3 Entrega fallida", "1.er fallo: reintento hoy; 2.º fallo: pedido «fallido» y escalado al supervisor"],
    ["4 Memoria larga", "Recupera el incidente de batería (similitud 0.79) y su estrategia de reasignación"],
    ["5 Cadena de eventos", "4 eventos en memoria, 3 replanes → aviso de inestabilidad"],
    ["6 LLM caído / 7 ambiguos / 8 plan imposible", "Degrada a reglas; pide datos o escala; 3 reparaciones y escalamiento"],
], [5.2, 12.6])
rich([("Hallazgo con el LLM local real (qwen2.5:1.5b): ", True),
      ("clasificó «Cliente ausente en el pedido 8» como inicio de jornada, lo que habría replanificado todo el "
       "día. Se añadieron guardarraíles (las reglas de alta precisión prevalecen; «nuevo_dia» solo con "
       "confirmación de reglas) y validación de los avisos redactados por el LLM, con pruebas de regresión.", False)])

# ------------------------------------------------------------------ 7
doc.add_heading("7. Decisiones de diseño y su justificación (IE6, IE9)", 1)
box("✍️ A completar por el equipo (sin apoyo de IA)", [
    "Para cada decisión, escriban su justificación técnica con sus palabras usando la evidencia indicada "
    "(detalle en docs/decisiones.md):",
    "• D1 LangGraph frente a otros frameworks de agentes — evidencia: flujo con ramas, estado persistente.",
    "• D2 Planificador determinista + LLM solo para interpretar/redactar — evidencia: hallazgo con qwen2.5:1.5b.",
    "• D3 Groq/Ollama intercambiables y 100% gratuito — evidencia: equipo con 7.9 GB de RAM y sin GPU dedicada.",
    "• D4 ChromaDB + nomic-embed-text y umbral 0.60 — evidencia: tabla de similitudes de la sección 4.",
    "[Escribir aquí]",
])

# ------------------------------------------------------------------ 8
doc.add_heading("8. Conclusiones y reflexiones individuales", 1)
box("✍️ A completar por cada integrante, sin apoyo de IA (obligatorio)", [
    "Integrante 1 — aprendizaje y contribución al proyecto: [escribir aquí]",
    "Integrante 2 — aprendizaje y contribución al proyecto: [escribir aquí]",
    "Conclusiones del equipo (logros, limitaciones y trabajo futuro): [escribir aquí]",
])

# ------------------------------------------------------------------ 9
doc.add_heading("9. Declaración de uso de inteligencia artificial", 1)
P("Se utilizó Claude Code (modelo Claude Sonnet 5.5, Anthropic) como asistente para generar el código "
  "base, las pruebas, los diagramas, el panel web y los borradores de la documentación técnica del repositorio. Todo el "
  "contenido fue ejecutado, revisado y validado por el equipo (57 pruebas automáticas y ejecución "
  "manual de los escenarios). Dentro del propio sistema, un LLM (Llama 3.3 vía Groq o Qwen2.5 vía Ollama) "
  "clasifica eventos y redacta avisos a clientes, con validación determinista de sus salidas. "
  "Las justificaciones técnicas, conclusiones y reflexiones individuales fueron redactadas por los "
  "integrantes sin apoyo de IA. [Equipo: confirmar o ajustar esta declaración.]")

# ------------------------------------------------------------------ 10
doc.add_heading("Referencias (APA 7)", 1)
refs = [
    "Chase, H. (2022). LangChain [Software]. LangChain, Inc. https://github.com/langchain-ai/langchain",
    "LangChain, Inc. (2024). LangGraph documentation. https://langchain-ai.github.io/langgraph/",
    "Chroma. (2023). Chroma: The AI-native open-source embedding database. https://www.trychroma.com/",
    "Groq, Inc. (2024). GroqCloud documentation. https://console.groq.com/docs",
    "Nussbaum, Z., Morris, J. X., Duderstadt, B., & Mulyar, A. (2024). Nomic Embed: Training a reproducible "
    "long context text embedder. arXiv:2402.01613.",
    "Ollama. (2023). Ollama [Software]. https://ollama.com",
    "Qwen Team. (2024). Qwen2.5 technical report. arXiv:2412.15115.",
    "Yao, S., Zhao, J., Yu, D., Du, N., Shafran, I., Narasimhan, K., & Cao, Y. (2023). ReAct: Synergizing "
    "reasoning and acting in language models. ICLR 2023. arXiv:2210.03629.",
]
for r in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(r)
    run.font.size = Pt(8.5)

out = OUT / "Informe_EP2_ISY0101.docx"
doc.save(out)
print("Informe guardado en", out)
